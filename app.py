# -*- coding: utf-8 -*-
"""语音及中枢大模型平台化功能清单 —— 按项目筛选工具（Streamlit 网页版）

运行方式：
    pip install streamlit openpyxl requests
    streamlit run app.py

流程：
    1. 上传 xlsx 文件，或输入在线文档链接（只读处理，不会修改原文档）；
       在线链接支持：可直接下载的 xlsx 链接、Google Sheets 链接
       （自动转换为官方 export 直链，需链接已设为“知道链接的任何人可查看”）；
    2. 自动读取“修订记录”页“项目”列的全部选项，支持多选项目（“所有项目”不默认勾选）；
    3. 点击“生成处理后的文件”，得到按所选项目筛选后的 xlsx 供下载。

处理规则：
    1) 删除“平台化计划”页；
    2) 保留“需求评审记录”“A2-基础语音回复”“腾讯视频（新）”页；
    3) “修订记录”及“附录6-权限管理-语音麦克风和定位权限”（项目纵向排列）
       只保留所选项目的条目；
    4) 其余每页：保留非项目列 + 所有选中具体项目对应的项目列，删除其他项目列；
       若某页所有列都不含任一选中具体项目，则删除该页。
"""
import io
import re

import openpyxl
import requests
import streamlit as st

from filter_core import get_project_options, normalize, process_bytes

# Google Sheets 链接 → 官方 export 直链（无需登录，可下载完整 xlsx）
GOOGLE_SHEET_RE = re.compile(r'docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)')


def normalize_url(url: str) -> str:
    """把 Google Sheets 分享链接改写为可直连下载的 export 链接；其余链接原样返回。"""
    m = GOOGLE_SHEET_RE.search(url)
    if m:
        return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"
    return url


@st.cache_data(show_spinner="正在从链接加载在线文档…")
def fetch_xlsx(url: str) -> bytes:
    """只读下载在线 xlsx 文档内容；绝不向原地址写回，不会修改在线文档。"""
    headers = {
        "User-Agent": ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                       "AppleWebKit/537.36 (KHTML, like Gecko) "
                       "Chrome/120.0 Safari/537.36"),
    }
    resp = requests.get(normalize_url(url), headers=headers, timeout=60)
    resp.raise_for_status()
    return resp.content


st.set_page_config(page_title="语音及中枢大模型功能清单筛选", layout="centered")
st.title("语音及中枢大模型平台化功能清单 — 按项目筛选")
st.caption("支持本地上传或在线文档链接，选择项目后一键生成筛选版（不修改原文档）。")

# ---------- 1. 数据来源：本地上传 / 在线文档链接 ----------
tab_up, tab_url = st.tabs(["📁 本地上传", "🔗 在线文档链接"])
with tab_up:
    uploaded = st.file_uploader("上传 xlsx 文件", type=["xlsx"])
with tab_url:
    url_input = st.text_input(
        "输入在线文档链接",
        placeholder="https://…（可下载的 xlsx 直链，或 Google Sheets 分享链接）",
    )
    st.caption("只读加载并处理后生成新文件，不会改动在线文档内容。"
               "Google Sheets 链接需已设置为“知道链接的任何人可查看”。")

raw, src_name = None, "功能清单"
if uploaded is not None:
    raw, src_name = uploaded.getvalue(), uploaded.name
elif url_input and url_input.strip():
    try:
        raw = fetch_xlsx(url_input.strip())
        src_name = url_input.strip().rsplit("/", 1)[-1] or "在线文档"
    except Exception as e:                        # noqa: BLE001
        st.error(f"在线文档加载失败：{e}（若是飞书/在线表格分享页，请先在本地导出为 xlsx 或改用本地上传）")
        raw = None

if raw is not None:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw))
    except Exception as e:                        # noqa: BLE001
        st.error(f"文档解析失败，请确认是有效的 xlsx 文件：{e}")
        st.stop()

    if "修订记录" not in wb.sheetnames:
        st.error("未在文件中找到“修订记录”页，无法读取项目选项。")
        st.stop()

    options = get_project_options(wb)
    if not options:
        st.error("“修订记录”页“项目”列中没有可选项。")
        st.stop()

    # ---------- 2. 多选项目（“所有项目”不默认勾选） ----------
    st.subheader("2. 选择要保留的项目（可多选）")
    st.caption(f"来源：{src_name}")
    selected = []
    n_cols = 3
    cols = st.columns(n_cols)
    for i, p in enumerate(options):
        if cols[i % n_cols].checkbox(p, key=f"ck_{i}_{normalize(p)}"):
            selected.append(p)

    if not selected:
        st.info("请至少勾选一个项目后再生成。")
    else:
        st.caption("已选择：" + "、".join(selected))
        if st.button("3. 生成处理后的文件", type="primary"):
            with st.spinner("正在处理，请稍候……"):
                processed = process_bytes(raw, selected)
            out_wb = openpyxl.load_workbook(io.BytesIO(processed), read_only=True)
            kept = out_wb.sheetnames
            deleted = [n for n in wb.sheetnames if n not in kept]
            st.success(f"处理完成：共保留 {len(kept)} 页，删除 {len(deleted)} 页。")
            if deleted:
                st.info("被删除的页：" + "、".join(deleted))
            fname = "语音及中枢大模型平台化功能清单-" + "-".join(selected)[:50] + ".xlsx"
            st.download_button(
                "下载处理后的 xlsx",
                data=processed,
                file_name=fname,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
