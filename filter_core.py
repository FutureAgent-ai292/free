# -*- coding: utf-8 -*-
"""语音及中枢大模型平台化功能清单 —— 按项目筛选核心处理逻辑（可独立测试）"""
import io
import re
import openpyxl

PROJ_RE = re.compile(r'^[A-Za-z]+\d')          # 项目列识别：表头以“字母+数字”开头
# 表头中提取项目代号的 token 规则（如 C62X-M17 / N66QB（对公）/ N50AB/N51AB / C66TB舱泊一体）
HEADER_CODE_RE = re.compile(r'[A-Z]{1,3}\d{1,4}[A-Z0-9]*(?:-[A-Z0-9]+)*')
KEEP_SHEETS = {'需求评审记录', 'A2-基础语音回复', '腾讯视频（新）'}   # 无条件保留的页
DEL_SHEETS  = {'平台化计划'}                     # 无条件删除的页
REV_SHEET   = '修订记录'                        # 按项目过滤条目的页
# 项目字段纵向排列（在“项目”列内按行存放）的页：按所选项目过滤行、保留全部列
ROW_FILTER_SHEETS = {'附录6-权限管理-语音麦克风和定位权限'}


def normalize(s: str) -> str:
    """去掉所有空白字符（含空格/换行/制表符），用于模糊匹配。"""
    return re.sub(r'\s+', '', str(s))


def get_project_options(wb: openpyxl.Workbook):
    """收集工作簿中出现的全部项目字段，供用户勾选：
        1) “修订记录”页“项目”列的全部取值（含“所有项目”）；
        2) 各功能页表头中的项目代号列（如 C62X-M17、B20CS、N66QB、B30X 等）。
    统一去空白归一化并去重，保持首次出现顺序，“所有项目”排在首位。
    """
    seen, order = set(), []

    def add(raw):
        key = normalize(raw)
        if key and key not in seen:
            seen.add(key)
            order.append(key)

    # 1) 修订记录 项目列
    if REV_SHEET in wb.sheetnames:
        ws = wb[REV_SHEET]
        hdr = [c.value for c in ws[1]]
        col = next((i for i, h in enumerate(hdr) if h and str(h).strip() == '项目'), None)
        if col is not None:
            for row in ws.iter_rows(min_row=2, values_only=True):
                v = row[col]
                if v is None:
                    continue
                for part in str(v).replace('\n', '|').split('|'):
                    if part.strip():
                        add(part.strip())

    # 2) 各功能页“项目列”的表头（提取项目代号 token，避开表头行中的备注文字）
    for name in wb.sheetnames:
        ws = wb[name]
        hrow = find_header_row(ws)
        if not hrow:
            continue
        for (c0, c1, header) in project_column_groups(ws, hrow):
            for tok in HEADER_CODE_RE.findall(normalize(str(header))):
                add(tok)

    # “所有项目”排首位
    if '所有项目' in order:
        order.remove('所有项目')
        order.insert(0, '所有项目')
    return order


def find_header_row(ws) -> int:
    """定位表头行（1 基）。

    判定规则：从上往下第一行满足 ①项目列>=2；或 ②仅 1 个项目列但整行非空单元格>=10
    （用于 A16/A22/A23/A24 这类只有单项目列的页，同时避免把“项目代号写在数据列
    （如附录6）”的数据行误判为表头）。
    """
    for i, row in enumerate(ws.iter_rows(max_row=10, values_only=True)):
        proj_n = sum(1 for v in row
                     if v is not None and str(v).strip() and PROJ_RE.match(str(v).strip()))
        nonempty = sum(1 for v in row if v is not None and str(v).strip())
        if proj_n >= 2 or (proj_n == 1 and nonempty >= 10):
            return i + 1
    return 0


def header_matches(header, chosen: str) -> bool:
    """判断表头是否代表所选项目（如 N53TB（C66TB-M05）含 C66TB-M05）。"""
    if not chosen:
        return False
    h, c = normalize(header), normalize(chosen)
    return bool(c) and (h == c or c in h)


def project_column_groups(ws, hrow: int):
    """返回 [(起始列, 结束列, 表头文本)]：每个项目列组（处理表头合并单元格）。"""
    hdr = [c.value for c in ws[hrow]]
    merges = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= hrow <= rng.max_row:
            for col in range(rng.min_col, rng.max_col + 1):
                if rng.min_row == hrow and rng.min_col != rng.max_col:
                    merges[col] = (rng.min_col, rng.max_col)   # 横向合并（如 N80KS 跨3列）
                else:
                    merges.setdefault(col, (col, col))
    groups = []
    for c in range(1, len(hdr) + 1):
        v = hdr[c - 1]
        if v is not None and str(v).strip() and PROJ_RE.match(str(v).strip()):
            c0, c1 = merges.get(c, (c, c))
            groups.append((c0, c1, v))
    return groups


def filter_rows_by_project(wb, sheet_name: str, selected_set):
    """按“项目”列过滤行：仅保留“项目”∈ 选中项目集合的条目（用于修订记录及纵向项目页）。"""
    ws = wb[sheet_name]
    hdr = [c.value for c in ws[1]]
    col = next((i for i, h in enumerate(hdr) if h and str(h).strip() == '项目'), None)
    if col is None:
        return
    norm_set = {normalize(s) for s in selected_set if s}
    if not norm_set:
        return
    for r in range(ws.max_row, 1, -1):
        v = ws.cell(row=r, column=col + 1).value
        keep = False
        if v is not None:
            nv = normalize(v)
            keep = any(s in nv for s in norm_set)
        if not keep:
            ws.delete_rows(r)


def process_workbook(wb, selected_projects: list):
    """按所选“项目”（可多选）处理工作簿（就地修改）。返回 wb。

    规则：
      1) 删除“平台化计划”页；
      2) 保留“需求评审记录”“A2-基础语音回复”“腾讯视频（新）”页；
      3) “修订记录”及纵向项目页（附录6）只保留“项目”∈ 所选项目的条目；
      4) 其余每页：保留非项目列 + 所有选中具体项目对应的项目列，删除其他项目列；
         若某页所有列都不含任一选中具体项目，则删除该页。
    """
    for name in DEL_SHEETS:
        if name in wb.sheetnames:
            del wb[name]

    specifics = [p for p in selected_projects if normalize(p) != '所有项目']

    # 行过滤类页：修订记录 + 纵向项目页
    for name in [REV_SHEET] + sorted(ROW_FILTER_SHEETS):
        if name in wb.sheetnames:
            filter_rows_by_project(wb, name, selected_projects)

    # 列过滤类页
    for name in list(wb.sheetnames):
        if name in KEEP_SHEETS or name == REV_SHEET or name in ROW_FILTER_SHEETS:
            continue
        ws = wb[name]
        hrow = find_header_row(ws)
        if hrow == 0 or not specifics:
            # 无项目列 / 未选具体项目：不满足“含所选具体项目列” → 删除该页
            if hrow == 0 and specifics:
                del wb[name]
            continue
        groups = project_column_groups(ws, hrow)
        keep_groups = [g for g in groups
                       if any(header_matches(g[2], s) for s in specifics)]
        if not keep_groups:
            del wb[name]                       # 所有列都不含任一选中项目 → 删页
            continue
        # 从右往左删除非所选项目列组，避免列号漂移；
        # 删除前先清除与待删列相交的合并单元格定义（openpyxl 的 delete_cols 不会清理它们）
        for c0, c1, _ in sorted((g for g in groups if g not in keep_groups),
                                key=lambda g: g[0], reverse=True):
            for rng in list(ws.merged_cells.ranges):
                if rng.min_col <= c1 and rng.max_col >= c0:
                    ws.unmerge_cells(str(rng))
            ws.delete_cols(c0, c1 - c0 + 1)
    return wb


def process_bytes(data: bytes, selected_projects: list) -> bytes:
    """从字节流读入、处理后输出字节流（供 Web 上传/在线链接使用）。"""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    process_workbook(wb, selected_projects)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
